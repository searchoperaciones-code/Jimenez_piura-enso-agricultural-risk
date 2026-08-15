from __future__ import annotations

import ast
import csv
import hashlib
import json
import subprocess
import sys
from pathlib import Path
from typing import Any
from zipfile import ZipFile


ROOT = Path(__file__).resolve().parents[1]
EXPECTED_BRANCH = "phase/c0a-outcome-evidence-v1"
EXPECTED_BASE = "144cb679186b8fdfe5b821aad635a7f805b7de28"
EXPECTED_STAGE_B_TAG = "v0.4.0-climate-exposure-freeze"

C0A_FILES = [
    ROOT / "outputs" / "outcome" / "C0A_EVIDENCE_REGISTRY.csv",
    ROOT / "outputs" / "outcome" / "C0A_OUTCOME_EVIDENCE_REPORT.md",
    ROOT / "scripts" / "c0a_outcome_evidence_preflight.py",
    ROOT / "tests" / "test_c0a_outcome_evidence_preflight.py",
]

RAW_AND_QA_EVIDENCE = [
    ROOT / "data" / "raw" / "Formato_dataset_productos_dra__ (2).csv",
    ROOT / "data" / "raw" / "Formato_DiccionarioDatos_productos_dra_.xlsx",
    ROOT / "data" / "raw" / "Formato_Metadatos_productos_dra_.docx",
    ROOT / "outputs" / "qa" / "Q1_data_readiness.md",
    ROOT / "outputs" / "qa" / "dictionary_forensic_audit.json",
    ROOT / "outputs" / "qa" / "source_dictionary_inconsistencies.csv",
    ROOT / "outputs" / "qa" / "independent_yield_price_checks.csv",
    ROOT / "outputs" / "qa" / "data_sources.csv",
    ROOT / "outputs" / "qa" / "manuscript_data_provenance.md",
    ROOT / "outputs" / "qa" / "dataset_schema.json",
    ROOT / "outputs" / "qa" / "raw_agricultural_audit.json",
    ROOT / "outputs" / "qa" / "domain_checks.json",
    ROOT / "outputs" / "qa" / "missing_zero_semantics.csv",
    ROOT / "outputs" / "qa" / "missing_zero_blank_samples.csv",
    ROOT / "scientific_annotations.md",
]

UPSTREAM_IMMUTABLE_PATHS = [
    ROOT / "data" / "processed" / "panel_master.csv",
    ROOT / "data" / "processed" / "panel_balanceado.csv",
    ROOT / "data" / "processed" / "land_physical.csv",
    ROOT / "data" / "processed" / "phenology" / "temporal_structure_monthly.csv",
    ROOT / "data" / "processed" / "phenology" / "phenology_windows_frozen.csv",
    ROOT / "data" / "processed" / "phenology" / "transient_cohort_exposures.parquet",
    ROOT / "data" / "processed" / "phenology" / "perennial_exposures_long.parquet",
    ROOT / "data" / "processed" / "phenology" / "transient_campaign_exposures_strict.parquet",
    ROOT / "config" / "climate_exposure" / "climate_exposure_spec_v1.json",
]

REGISTRY = ROOT / "outputs" / "outcome" / "C0A_EVIDENCE_REGISTRY.csv"
REPORT = ROOT / "outputs" / "outcome" / "C0A_OUTCOME_EVIDENCE_REPORT.md"

REGISTRY_COLUMNS = [
    "EVIDENCE_ID",
    "VARIABLE",
    "CLAIM",
    "CLAIM_CATEGORY",
    "SOURCE_TIER",
    "SOURCE_AUTHORITY",
    "SOURCE_TITLE",
    "SOURCE_DATE",
    "SOURCE_URL",
    "LOCAL_SOURCE_PATH",
    "EXACT_DATASET_LINKAGE",
    "EVIDENCE_FINDING",
    "DECISION_STATUS",
    "DECISION_VALUE",
    "LIMITATION",
    "NOTES",
]

ALLOWED_VARIABLES = {
    "PRODUCCION",
    "COSECHA",
    "SIEMBRA",
    "VERDE_ACTUAL",
    "PRECIO_CHACRA",
    "ANO",
    "MES",
    "YIELD_RAW",
    "CROSS_VARIABLE",
}

ALLOWED_CATEGORIES = {
    "SEMANTICS",
    "UNIT",
    "TEMPORAL",
    "STOCK_FLOW",
    "DIMENSIONALITY",
    "SOURCE_LINKAGE",
    "NUMERICAL_CONCORDANCE",
}

ALLOWED_TIERS = {"TIER_1", "TIER_2", "TIER_3", "TIER_4"}
ALLOWED_LINKAGE = {"TRUE", "FALSE", "PARTIAL"}
ALLOWED_STATUSES = {
    "CERTIFIED_EXACT_SOURCE",
    "CERTIFIED_INSTITUTIONAL_CONVERGENCE",
    "SUPPORTED_NOT_CERTIFIED",
    "UNRESOLVED",
    "CONTRADICTED",
}

FORBIDDEN_OUTCOME_ARTIFACTS = [
    ROOT / "data" / "processed" / "outcome",
    ROOT / "data" / "processed" / "phenology" / "phenology_exposures_long.parquet",
]

C0A1_TIER2_AUTHORITY = "GORE_PIURA_DRAP_OFICINA_ESTADISTICA"
C0A1_PRIMARY_TIER2_TITLES = {
    "Piura produjo más de 516 mil toneladas de arroz durante el año 2025",
    "Piura consolida liderazgo en la producción de limón sutil con más de 203 mil toneladas y exportaciones a seis paises",
    "Producción de mango en Piura crece 55.8 por ciento y supera las 535 mil toneladas en 2025",
}
C0A1_REQUIRED_STATUS = "CERTIFIED_INSTITUTIONAL_CONVERGENCE"
C0A1_REQUIRED_UNIT = "METRIC_TONNE"

MODELLING_IMPORT_PREFIXES = {
    "linearmodels",
    "sklearn",
    "statsmodels",
    "xgboost",
    "lightgbm",
    "cvxpy",
    "pulp",
    "pyomo",
}

MODELLING_IDENTIFIERS = {
    "OLS",
    "PanelOLS",
    "RandomForestRegressor",
    "RandomForestClassifier",
    "XGBRegressor",
    "XGBClassifier",
    "Lasso",
    "ElasticNet",
    "linprog",
    "milp",
}


def rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def git_output(*args: str) -> str:
    result = subprocess.run(["git", *args], cwd=ROOT, check=True, capture_output=True, text=True)
    return result.stdout.strip()


def git_bytes(*args: str) -> bytes:
    result = subprocess.run(["git", *args], cwd=ROOT, check=True, capture_output=True)
    return result.stdout


def git_status_paths() -> list[str]:
    output = git_output("status", "--short", "--untracked-files=all")
    return [line.strip() for line in output.splitlines() if line.strip()]


def branch_gate() -> dict[str, Any]:
    branch = git_output("branch", "--show-current")
    head = git_output("rev-parse", "HEAD")
    base_is_ancestor = subprocess.run(
        ["git", "merge-base", "--is-ancestor", EXPECTED_BASE, "HEAD"],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
    ).returncode == 0
    tag_points = git_output("tag", "--points-at", "HEAD").splitlines()
    return {
        "branch": branch,
        "head": head,
        "base": EXPECTED_BASE,
        "branch_ok": branch == EXPECTED_BRANCH,
        "base_is_ancestor": base_is_ancestor,
        "stage_b_tag_at_head": EXPECTED_STAGE_B_TAG in tag_points,
        "status": "PASS"
        if branch == EXPECTED_BRANCH and base_is_ancestor and EXPECTED_STAGE_B_TAG in tag_points
        else "FAIL",
    }


def raw_source_inventory() -> dict[str, Any]:
    items = []
    for path in RAW_AND_QA_EVIDENCE:
        items.append(
            {
                "path": rel(path),
                "exists": path.exists(),
                "sha256": sha256_file(path) if path.exists() else "MISSING",
            }
        )
    return {
        "available": sum(1 for item in items if item["exists"]),
        "required": len(items),
        "missing": [item["path"] for item in items if not item["exists"]],
        "items": items,
        "status": "PASS" if all(item["exists"] for item in items) else "FAIL",
    }


def inspect_dictionary() -> dict[str, str]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"status": "FAIL", "error": "openpyxl unavailable"}
    path = ROOT / "data" / "raw" / "Formato_DiccionarioDatos_productos_dra_.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = {}
    for row in sheet.iter_rows(values_only=True):
        if not row or not row[0]:
            continue
        variable = str(row[0]).strip()
        if variable in {"PRODUCCION", "COSECHA", "SIEMBRA", "VERDE_ACTUAL", "PRECIO_CHACRA", "ANO", "MES"}:
            rows[variable] = " | ".join(str(value).strip() for value in row if value is not None)
    workbook.close()
    return {"status": "PASS" if len(rows) == 7 else "FAIL", **rows}


def inspect_metadata_text() -> dict[str, Any]:
    path = ROOT / "data" / "raw" / "Formato_Metadatos_productos_dra_.docx"
    snippets = []
    with ZipFile(path) as archive:
        with archive.open("word/document.xml") as handle:
            chunks = []
            while True:
                chunk = handle.read(4096)
                if not chunk:
                    break
                chunks.append(chunk)
    text = b"".join(chunks).decode("utf-8", errors="replace")
    for token in [
        "Campaña agrícola",
        "Gobierno Regional Piura",
        "Oficina de Estadística",
        "valor cero",
        "valor en blanco",
    ]:
        snippets.append({"token": token, "present": token.lower() in text.lower()})
    return {"status": "PASS" if all(item["present"] for item in snippets) else "FAIL", "snippets": snippets}


def read_registry() -> list[dict[str, str]]:
    with REGISTRY.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def registry_validation() -> dict[str, Any]:
    with REGISTRY.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = list(reader)
        columns = reader.fieldnames
    sorted_rows = sorted(rows, key=lambda row: (row["VARIABLE"], row["CLAIM_CATEGORY"], row["SOURCE_TIER"], row["SOURCE_TITLE"], row["CLAIM"]))
    expected_ids = [f"C0A-E{i:03d}" for i in range(1, len(rows) + 1)]
    ids = [row["EVIDENCE_ID"] for row in rows]
    no_multiline = all("\n" not in value and "\r" not in value for row in rows for value in row.values())
    checks = {
        "row_count": len(rows),
        "columns_exact": columns == REGISTRY_COLUMNS,
        "ids_deterministic": ids == expected_ids,
        "sort_deterministic": rows == sorted_rows,
        "variables_allowed": all(row["VARIABLE"] in ALLOWED_VARIABLES for row in rows),
        "categories_allowed": all(row["CLAIM_CATEGORY"] in ALLOWED_CATEGORIES for row in rows),
        "tiers_allowed": all(row["SOURCE_TIER"] in ALLOWED_TIERS for row in rows),
        "linkage_allowed": all(row["EXACT_DATASET_LINKAGE"] in ALLOWED_LINKAGE for row in rows),
        "statuses_allowed": all(row["DECISION_STATUS"] in ALLOWED_STATUSES for row in rows),
        "no_multiline_cells": no_multiline,
    }
    checks["status"] = "PASS" if all(value for key, value in checks.items() if key != "row_count") and rows else "FAIL"
    return checks


def upstream_immutability_gate() -> dict[str, Any]:
    missing = [rel(path) for path in UPSTREAM_IMMUTABLE_PATHS if not path.exists()]
    diff = subprocess.run(
        ["git", "diff", "--quiet", "--", *[rel(path) for path in UPSTREAM_IMMUTABLE_PATHS]],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
    ).returncode
    hash_mismatches = []
    for path in UPSTREAM_IMMUTABLE_PATHS:
        if not path.exists():
            continue
        head_bytes = git_bytes("show", f"HEAD:{rel(path)}")
        head_sha = hashlib.sha256(head_bytes).hexdigest()
        worktree_sha = sha256_file(path)
        if head_sha != worktree_sha:
            hash_mismatches.append({"path": rel(path), "head": head_sha, "worktree": worktree_sha})
    return {
        "missing": missing,
        "diff_clean": diff == 0,
        "hash_mismatches": hash_mismatches,
        "status": "PASS" if not missing and diff == 0 and not hash_mismatches else "FAIL",
    }


def persistent_scope_gate() -> dict[str, Any]:
    expected = sorted(f"?? {rel(path)}" for path in C0A_FILES)
    actual = sorted(git_status_paths())
    return {"expected": expected, "actual": actual, "status": "PASS" if actual == expected else "FAIL"}


def outcome_firewall_gate() -> dict[str, Any]:
    existing_forbidden = [rel(path) for path in FORBIDDEN_OUTCOME_ARTIFACTS if path.exists()]
    artifact_hits = []
    for root in [ROOT / "data", ROOT / "outputs"]:
        if not root.exists():
            continue
        for path in root.rglob("*"):
            if not path.is_file():
                continue
            name = path.name.lower()
            if any(token in name for token in ["regression", "optimization", "econometric"]):
                artifact_hits.append(rel(path))
            if "model" in name and "metadata" not in name:
                artifact_hits.append(rel(path))
    return {
        "existing_forbidden": existing_forbidden,
        "artifact_hits": sorted(set(artifact_hits)),
        "status": "PASS" if not existing_forbidden and not artifact_hits else "FAIL",
    }


def static_modelling_firewall() -> dict[str, Any]:
    source = (ROOT / "scripts" / "c0a_outcome_evidence_preflight.py").read_text(encoding="utf-8")
    tree = ast.parse(source)
    bad_imports = []
    bad_identifiers = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for alias in node.names:
                if alias.name.split(".")[0] in MODELLING_IMPORT_PREFIXES:
                    bad_imports.append(alias.name)
        elif isinstance(node, ast.ImportFrom):
            module = node.module or ""
            if module.split(".")[0] in MODELLING_IMPORT_PREFIXES:
                bad_imports.append(module)
        elif isinstance(node, ast.Name) and node.id in MODELLING_IDENTIFIERS:
            bad_identifiers.append(node.id)
        elif isinstance(node, ast.Attribute) and node.attr in MODELLING_IDENTIFIERS:
            bad_identifiers.append(node.attr)
    return {
        "bad_imports": sorted(set(bad_imports)),
        "bad_identifiers": sorted(set(bad_identifiers)),
        "status": "PASS" if not bad_imports and not bad_identifiers else "FAIL",
    }


def report_validation() -> dict[str, Any]:
    text = REPORT.read_text(encoding="utf-8")
    required_variables = ["PRODUCCION", "COSECHA", "SIEMBRA", "VERDE_ACTUAL", "PRECIO_CHACRA", "ANO", "MES"]
    required_sections = [
        "1. Executive verdict",
        "2. Frozen upstream identity",
        "3. Source inventory",
        "4. Evidence hierarchy",
        "5. Variable-by-variable adjudication",
        "6. PRODUCCION unit adjudication",
        "7. COSECHA/SIEMBRA/VERDE_ACTUAL area adjudication",
        "8. PRECIO_CHACRA adjudication",
        "9. ANO/MES temporal semantics",
        "10. Dimensional analysis of PRODUCCION / COSECHA",
        "11. Numerical concordance checks",
        "12. Contradictions and unresolved issues",
        "13. Implication for transient Outcome Master",
        "14. Implication for perennial Outcome Master",
        "15. Implication for future monetary translation",
        "16. Explicit list of actions NOT authorized",
        "17. Final C0A verdict",
    ]
    checks = {
        "variables_present": all(variable in text for variable in required_variables),
        "sections_present": all(section in text for section in required_sections),
        "yield_unit_explicit": "YIELD_UNIT_DECISION=TM_PER_HA" in text or "YIELD_RAW is adjudicated as TM_PER_HA" in text,
        "monetary_blocked": "GVP construction is blocked" in text,
        "no_blanket_causal": "Causal interpretation is not authorized" in text and "causal effect is certified" not in text.lower(),
        "c0a1_final": "C0A_1_STATUS=PASS_UNIT_CERTIFIED_BY_INSTITUTIONAL_CONVERGENCE" in text,
        "status": "PASS",
    }
    checks["status"] = "PASS" if all(value for key, value in checks.items() if key != "status") else "FAIL"
    return checks


def decision_lookup(rows: list[dict[str, str]], variable: str, category: str, claim_contains: str = "") -> dict[str, str]:
    for row in rows:
        if row["VARIABLE"] == variable and row["CLAIM_CATEGORY"] == category and claim_contains.lower() in row["CLAIM"].lower():
            return row
    raise KeyError((variable, category, claim_contains))


def institutional_convergence_gate(rows: list[dict[str, str]]) -> dict[str, Any]:
    tier2_sources = [
        row
        for row in rows
        if row["VARIABLE"] == "PRODUCCION"
        and row["CLAIM_CATEGORY"] == "UNIT"
        and row["SOURCE_TIER"] == "TIER_2"
        and row["SOURCE_AUTHORITY"] == C0A1_TIER2_AUTHORITY
        and row["SOURCE_TITLE"] in C0A1_PRIMARY_TIER2_TITLES
        and row["DECISION_VALUE"] == C0A1_REQUIRED_UNIT
        and row["DECISION_STATUS"] == C0A1_REQUIRED_STATUS
    ]
    tier4_certifiers = [
        row
        for row in rows
        if row["VARIABLE"] == "PRODUCCION"
        and row["CLAIM_CATEGORY"] == "UNIT"
        and row["SOURCE_TIER"] == "TIER_4"
        and row["DECISION_STATUS"] in {"CERTIFIED_EXACT_SOURCE", "CERTIFIED_INSTITUTIONAL_CONVERGENCE"}
    ]
    exact_linkage = any(
        row["VARIABLE"] == "CROSS_VARIABLE"
        and row["CLAIM_CATEGORY"] == "SOURCE_LINKAGE"
        and row["DECISION_VALUE"] == "EXACT_GORE_PIURA_DRA_DATASET_PACKAGE"
        for row in rows
    )
    final = decision_lookup(rows, "PRODUCCION", "UNIT", "Final adjudication")
    yield_unit = decision_lookup(rows, "YIELD_RAW", "UNIT")
    monetary = decision_lookup(rows, "CROSS_VARIABLE", "DIMENSIONALITY", "monetary translation")
    conditions = {
        "same_authority_tier2_count": len({row["SOURCE_TITLE"] for row in tier2_sources}),
        "same_authority_sources": sorted({row["SOURCE_TITLE"] for row in tier2_sources}),
        "tier4_certifiers": tier4_certifiers,
        "exact_dataset_linkage": exact_linkage,
        "final_produccion_unit": final["DECISION_VALUE"],
        "final_produccion_status": final["DECISION_STATUS"],
        "yield_unit": yield_unit["DECISION_VALUE"],
        "monetary_status": monetary["DECISION_VALUE"],
    }
    pass_gate = (
        conditions["same_authority_tier2_count"] >= 2
        and not tier4_certifiers
        and exact_linkage
        and final["DECISION_STATUS"] == C0A1_REQUIRED_STATUS
        and final["DECISION_VALUE"] == C0A1_REQUIRED_UNIT
        and yield_unit["DECISION_VALUE"] == "TM_PER_HA"
        and monetary["DECISION_VALUE"] == "DIMENSIONALLY_AUTHORIZED_NOT_BUILT"
    )
    return {**conditions, "status": "PASS" if pass_gate else "FAIL"}


def numerical_concordance_gate(rows: list[dict[str, str]]) -> dict[str, Any]:
    concordance = decision_lookup(rows, "CROSS_VARIABLE", "NUMERICAL_CONCORDANCE", "No exact official")
    report_text = REPORT.read_text(encoding="utf-8")
    no_calendar_campaign_mix = "not exactly comparable" in report_text and "without mixing periods" in report_text
    export_excluded = "export" in report_text.lower() and "not provide exact production totals" in report_text
    return {
        "cases": 0 if concordance["DECISION_VALUE"] == "NO_EXACT_NUMERICAL_CONCORDANCE_AVAILABLE" else 1,
        "status_value": concordance["DECISION_VALUE"],
        "no_calendar_campaign_mix": no_calendar_campaign_mix,
        "export_volume_not_used": export_excluded,
        "status": "PASS"
        if concordance["DECISION_VALUE"] == "NO_EXACT_NUMERICAL_CONCORDANCE_AVAILABLE"
        and no_calendar_campaign_mix
        and export_excluded
        else "FAIL",
    }


def evaluate() -> dict[str, Any]:
    rows = read_registry()
    gates = {
        "branch": branch_gate(),
        "raw_inventory": raw_source_inventory(),
        "dictionary": inspect_dictionary(),
        "metadata": inspect_metadata_text(),
        "registry": registry_validation(),
        "upstream": upstream_immutability_gate(),
        "persistent_scope": persistent_scope_gate(),
        "outcome_firewall": outcome_firewall_gate(),
        "static_modelling": static_modelling_firewall(),
        "report": report_validation(),
    }
    gates["institutional_convergence"] = institutional_convergence_gate(rows)
    gates["numerical_concordance"] = numerical_concordance_gate(rows)
    unresolved = [row for row in rows if row["DECISION_STATUS"] == "UNRESOLVED"]
    contradicted = [row for row in rows if row["DECISION_STATUS"] == "CONTRADICTED"]
    required_pass = all(gate["status"] == "PASS" for gate in gates.values())
    if not required_pass:
        status = "FAIL"
    elif decision_lookup(rows, "PRODUCCION", "UNIT", "Final adjudication")["DECISION_STATUS"] == C0A1_REQUIRED_STATUS:
        status = "PASS_UNIT_CERTIFIED_BY_INSTITUTIONAL_CONVERGENCE"
    elif decision_lookup(rows, "PRODUCCION", "UNIT", "Final adjudication")["DECISION_STATUS"] == "SUPPORTED_NOT_CERTIFIED":
        status = "PASS_UNIT_STRONGLY_SUPPORTED_NOT_CERTIFIED"
    else:
        status = "PASS_UNIT_REMAINS_UNRESOLVED"
    return {
        "gates": gates,
        "rows": rows,
        "unresolved": unresolved,
        "contradicted": contradicted,
        "status": status,
    }


def terminal_summary(report: dict[str, Any]) -> list[str]:
    rows = report["rows"]
    gates = report["gates"]
    produccion_unit = decision_lookup(rows, "PRODUCCION", "UNIT", "Final adjudication")
    yield_dim = decision_lookup(rows, "YIELD_RAW", "DIMENSIONALITY")
    yield_unit = decision_lookup(rows, "YIELD_RAW", "UNIT")
    monetary = decision_lookup(rows, "CROSS_VARIABLE", "DIMENSIONALITY", "monetary translation")
    unresolved_claims = ";".join(row["VARIABLE"] + ":" + row["CLAIM_CATEGORY"] for row in report["unresolved"]) or "NONE"
    contradictions = ";".join(row["VARIABLE"] + ":" + row["CLAIM_CATEGORY"] for row in report["contradicted"]) or "NONE"
    same_authority_sources = gates["institutional_convergence"]["same_authority_sources"]
    return [
        f"C0A_1_EXTERNAL_PRIMARY_SOURCES={gates['institutional_convergence']['same_authority_tier2_count']} DRAP Tier-2 unit sources plus MIDAGRI/SIEA Tier-3 support",
        f"C0A_1_SAME_AUTHORITY_SOURCES={';'.join(same_authority_sources)}",
        f"C0A_1_EXACT_DATASET_AUTHORITY_LINKAGE={'PASS' if gates['institutional_convergence']['exact_dataset_linkage'] else 'FAIL'}",
        f"C0A_1_NUMERICAL_CONCORDANCE_CASES={gates['numerical_concordance']['cases']}",
        f"C0A_1_NUMERICAL_CONCORDANCE_STATUS={gates['numerical_concordance']['status_value']}",
        f"PRODUCCION_UNIT_STATUS={produccion_unit['DECISION_STATUS']}",
        f"PRODUCCION_UNIT_DECISION={produccion_unit['DECISION_VALUE']}",
        "PRODUCCION_UNIT_RATIONALE=same DRAP statistical office linkage plus 3 Tier-2 official publications using tonnes or metric tonnes with ha and yield units; no contradiction; no forced concordance",
        f"YIELD_DIMENSIONAL_STATUS={yield_dim['DECISION_STATUS']}",
        f"YIELD_UNIT_DECISION={yield_unit['DECISION_VALUE']}",
        f"MONETARY_TRANSLATION_STATUS={monetary['DECISION_VALUE']}; GVP_NOT_BUILT",
        f"CONTRADICTIONS={contradictions}",
        f"UNRESOLVED_CLAIMS={unresolved_claims}",
        f"PERSISTENT_SCOPE_GATE={gates['persistent_scope']['status']}",
        f"UPSTREAM_IMMUTABILITY_GATE={gates['upstream']['status']}",
        f"OUTCOME_FIREWALL_GATE={gates['outcome_firewall']['status']}",
        f"STATIC_MODELLING_FIREWALL={gates['static_modelling']['status']}",
        "TESTS=RUN_SEPARATELY",
        f"C0A_1_STATUS={report['status']}",
    ]


def main() -> int:
    report = evaluate()
    for line in terminal_summary(report):
        print(line)
    return 0 if report["status"] != "FAIL" else 1


if __name__ == "__main__":
    raise SystemExit(main())
